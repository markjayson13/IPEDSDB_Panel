#!/usr/bin/env python3
"""Build and verify a repair's analyst exports, then bind a publication receipt.

Read the source panel and canonical lineage once for the labeled Parquet seed.
CSV, Stata, and Excel exports then use its embedded definitions and provenance.
All outputs are written only beneath the explicitly supplied output directory.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import traceback

import duckdb
import pandas as pd
import pyarrow as pa
import pyarrow.csv as pcsv
import pyarrow.dataset as ds
import pyarrow.parquet as pq
from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1]
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from panel_export import sha256_file
from source_metadata_corrections import DEFAULT_REGISTRY, load_registry, normalize_number


def run_checked(command: list[str], log: Path, *, cwd: Path = SCRIPTS.parent,
                native_output: bool = False) -> None:
    # macOS Stata uses inherited PWD to place its batch log, even when the
    # process working directory was changed by subprocess.
    environment = {**os.environ, "PWD": str(cwd.resolve())} if native_output else None
    result = subprocess.run(command, cwd=cwd, env=environment, capture_output=True, text=True)
    output = result.stdout + result.stderr
    if native_output and any(marker in output for marker in ("Stata license:", "Serial number:", "Licensed to:")):
        start = output.find("\n. do ")
        output = output[start + 1:] if start >= 0 else "Application banner omitted; no validation command output captured.\n"
    log.write_text(output, encoding="utf-8")
    if result.returncode:
        raise RuntimeError(f"Command failed with status {result.returncode}; see {log}: {output[-2500:]}")


def affected_columns(schema: pa.Schema, mapping_path: Path, registry: dict, year: int) -> list[str]:
    required = {(row["varname"].upper(), normalize_number(row["varnumber"]))
                for row in registry["corrections"] if int(row["year"]) == year}
    if not required:
        raise ValueError(f"The correction registry has no variables for {year}")
    if mapping_path.suffix.lower() == ".csv":
        with mapping_path.open(newline="", encoding="utf-8-sig") as handle:
            mapping = list(csv.DictReader(handle))
    else:
        fields = set(pq.read_schema(mapping_path).names)
        output_field = next((name for name in ("analysis_column", "output_column") if name in fields), None)
        number_field = next((name for name in ("source_varnumber", "varnumber") if name in fields), None)
        if not output_field or not number_field or "varname" not in fields:
            raise ValueError("Mapping requires an output column, varname, and source variable number")
        projection = f'"{output_field}" AS analysis_column, "{number_field}" AS source_varnumber, varname'
        predicate = f'upper(CAST(varname AS VARCHAR)) IN ({", ".join("?" for _ in required)})'
        parameters = [str(mapping_path), *sorted(identity[0] for identity in required)]
        if "year" in fields:
            projection += ", year"
            predicate += ' AND try_cast(year AS BIGINT) = ?'
            parameters.append(year)
        with tempfile.TemporaryDirectory(prefix="ipeds-repair-mapping-") as spill:
            with duckdb.connect(config={"memory_limit": "256MB", "temp_directory": spill, "threads": "2"}) as connection:
                query = connection.execute(f'SELECT DISTINCT {projection} FROM read_parquet(?) WHERE {predicate}', parameters)
                fetch_table = getattr(query, "to_arrow_table", query.fetch_arrow_table)
                mapping = fetch_table().to_pylist()
    matches = {identity: set() for identity in required}
    for row in mapping:
        if row.get("year") not in (None, "") and int(row["year"]) != year:
            continue
        identity = (str(row.get("varname", "")).upper(),
                    normalize_number(row.get("source_varnumber", row.get("varnumber", ""))))
        output = row.get("analysis_column", row.get("output_column"))
        if identity in matches and output in schema.names:
            matches[identity].add(output)
    unresolved = {str(identity): sorted(columns) for identity, columns in matches.items() if len(columns) != 1}
    if unresolved:
        raise ValueError(f"Correction identities require exactly one physical panel column: {unresolved}")
    chosen = {next(iter(columns)) for columns in matches.values()}
    if len(chosen) != len(required):
        raise ValueError("Multiple correction identities collapse to one output; this validator requires direct corrected variables.")
    return sorted(chosen)


def read_metadata(path: Path) -> dict:
    metadata = json.loads(Path(str(path) + ".metadata.json").read_text(encoding="utf-8"))
    if metadata["data_sha256"] != sha256_file(path):
        raise ValueError(f"Companion checksum does not match {path}")
    if metadata.get("readiness_status", metadata["metadata_status"]) != "complete":
        raise ValueError(f"Export is not ready: {path}")
    return metadata


def verify_representation(path: Path, expected: pa.Table) -> dict:
    metadata = read_metadata(path)
    fmt = path.suffix.lower()
    result = {"path": str(path.resolve()), "metadata_status": metadata["metadata_status"],
              "readiness_status": metadata.get("readiness_status"),
              "comparability_status": metadata.get("comparability_status"), "rows": expected.num_rows}
    if fmt == ".parquet":
        actual = pq.read_table(path)
        if not actual.equals(expected, check_metadata=False):
            raise ValueError("Parquet observations differ from the source projection")
        if not all((field.metadata or {}).get(b"ipeds:variable") for field in actual.schema):
            raise ValueError("Parquet field metadata is missing")
        result["embedded_field_metadata"] = True
    elif fmt == ".csv":
        actual = pcsv.read_csv(path, convert_options=pcsv.ConvertOptions(
            column_types={field.name: field.type for field in expected.schema}, null_values=[""],
            strings_can_be_null=True, quoted_strings_can_be_null=False))
        if not actual.equals(expected, check_metadata=False):
            raise ValueError("CSV observations differ from the source projection")
    elif fmt == ".dta":
        with pd.read_stata(path, iterator=True, convert_categoricals=False) as reader:
            labels, values = reader.variable_labels(), reader.value_labels()
            actual = reader.read()
        mapping = {row["name"]: row["export_name"] for row in metadata["variables"]}
        reference = expected.to_pandas().rename(columns=mapping)
        pd.testing.assert_frame_equal(actual, reference, check_dtype=False, check_exact=True)
        for variable in metadata["variables"]:
            name = variable["export_name"]
            if labels[name] != variable["export_label"]:
                raise ValueError(f"Native Stata variable label mismatch: {name}")
            expected_codes = {int(key): value for key, value in variable["stata_value_labels"].items()}
            if values.get(name, {}) != expected_codes:
                raise ValueError(f"Native Stata value labels mismatch: {name}")
        result.update(variable_labels_checked=len(labels), value_label_sets_checked=len(values))
    elif fmt == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            data = workbook["data"]
            rows = list(data.values)
            actual = pd.DataFrame(rows[1:], columns=rows[0])
            pd.testing.assert_frame_equal(actual, expected.to_pandas(), check_dtype=False, check_exact=True)
            for row in data.iter_rows(min_row=2):
                for field, cell in zip(expected.schema, row):
                    if cell.value is not None and (pa.types.is_integer(field.type) or pa.types.is_floating(field.type)) and cell.data_type != "n":
                        raise ValueError(f"Excel numeric observation became a nonnumeric cell: {field.name}")
            dictionary = list(workbook["dictionary"].values)
            if "comparability_status" not in dictionary[0] or "units_status" not in dictionary[0]:
                raise ValueError("Excel dictionary omits measurement-status metadata")
            result["sheets"] = workbook.sheetnames
        finally:
            workbook.close()
    else:
        raise ValueError(f"Unsupported verification format: {fmt}")
    result["exact_observation_roundtrip"] = True
    return result


def _stata_string(value: str) -> str:
    # Encode source labels as codepoints, never as executable do-file text.
    return "+".join(f"uchar({ord(character)})" for character in value) or '""'


def sanitize_native_log(log: Path) -> str:
    """Retain validation commands/results without the application license banner."""
    text = log.read_text(errors="replace")
    command_start = text.find("\n. do ")
    if command_start < 0:
        raise ValueError("Native log has no do-file execution record")
    text = text[command_start + 1:]
    log.write_text(text, encoding="utf-8")
    return text


def verify_native_stata(path: Path, expected: pa.Table, binary: Path, proof_dir: Path) -> dict:
    if not binary.is_file():
        raise ValueError(f"Native Stata executable is unavailable: {binary}")
    metadata = read_metadata(path)
    aliases = {row["name"]: row["export_name"] for row in metadata["variables"]}
    names = {name.upper(): name for name in expected.column_names}
    keys = [names["YEAR"], names["UNITID"]]
    reference_names = {name: aliases[name] if name in keys else f"r{index:04d}"
                       for index, name in enumerate(expected.column_names)}
    reference_csv = proof_dir / "native_reference.csv"
    pcsv.write_csv(expected.rename_columns([reference_names[name] for name in expected.column_names]), reference_csv)
    numeric_positions = [str(index + 1) for index, field in enumerate(expected.schema)
                         if pa.types.is_integer(field.type) or pa.types.is_floating(field.type) or pa.types.is_boolean(field.type)]
    string_positions = [str(index + 1) for index, field in enumerate(expected.schema)
                        if pa.types.is_string(field.type) or pa.types.is_large_string(field.type)]
    options = " asdouble"
    if numeric_positions:
        options += " numericcols(" + " ".join(numeric_positions) + ")"
    if string_positions:
        options += " stringcols(" + " ".join(string_positions) + ")"
    key_names = " ".join(aliases[key] for key in keys)
    commands = ["clear all", "set more off",
                f'import delimited using "{reference_csv}", clear varnames(1) case(preserve) encoding("utf-8"){options}',
                "tempfile reference", 'save "`reference\'", replace',
                f'use "{path.resolve()}", clear', f"assert _N == {expected.num_rows}", f"isid {key_names}"]
    for variable in metadata["variables"]:
        name = variable["export_name"]
        commands.append(f'mata: assert(st_varlabel("{name}") == {_stata_string(variable["export_label"])})')
        for code, label in variable.get("stata_value_labels", {}).items():
            commands.append(f'mata: assert(st_vlmap(st_varvaluelabel("{name}"), {int(code)}) == {_stata_string(label)})')
    commands += [f'merge 1:1 {key_names} using "`reference\'"', "assert _merge == 3"]
    for name in expected.column_names:
        if name not in keys:
            commands.append(f"assert {aliases[name]} == {reference_names[name]}")
    commands += ['display "NATIVE_STATA_REPAIR_EXPORT_PASS"', "exit, clear"]
    do_file = proof_dir / "native_verify.do"
    do_file.write_text("\n".join(commands) + "\n", encoding="utf-8")
    log = proof_dir / "native_verify.log"
    if log.exists():
        raise ValueError("Native validation log already exists; use a fresh output directory")
    run_checked([str(binary), "-e", "do", str(do_file)], proof_dir / "native_process.txt", cwd=proof_dir, native_output=True)
    if not log.is_file() or "NATIVE_STATA_REPAIR_EXPORT_PASS" not in sanitize_native_log(log).splitlines():
        raise ValueError(f"Native Stata did not produce a passing receipt; inspect {proof_dir}")
    return {"status": "pass", "executable": str(binary), "log": str(log.resolve()),
            "log_sha256": sha256_file(log), "rows_compared": expected.num_rows,
            "columns_compared": expected.num_columns, "variable_labels_checked": expected.num_columns,
            "value_labels_checked": sum(len(row.get("stata_value_labels", {})) for row in metadata["variables"])}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    for argument in ("input", "dictionary", "codes", "lineage", "variable-map", "output-dir"):
        parser.add_argument("--" + argument, required=True, type=Path)
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY)
    parser.add_argument("--year", type=int, default=2023)
    parser.add_argument("--stata-binary", type=Path, default=Path("/Applications/StataNow/StataBE.app/Contents/MacOS/StataBE"))
    args = parser.parse_args()
    output = args.output_dir.resolve()
    if output.exists() and any(output.iterdir()):
        raise ValueError("Validation output directory must be new or empty; receipts must not reuse stale artifacts")
    output.mkdir(parents=True, exist_ok=True)
    proof = output / "_validation"
    proof.mkdir()
    receipt = {"status": "fail", "created_utc": datetime.now(timezone.utc).isoformat(),
               "validator": str(Path(__file__).resolve()), "validator_sha256": sha256_file(Path(__file__)),
               "year": args.year, "formats": {}, "artifacts": []}
    try:
        registry = load_registry(args.registry)
        dataset = ds.dataset(args.input, format="parquet")
        columns = affected_columns(dataset.schema, args.variable_map, registry, args.year)
        variables_file = proof / "selected_variables.txt"
        variables_file.write_text("\n".join(columns) + "\n")
        receipt["corrected_variable_count"] = len(columns)
        seed = output / f"affected_{args.year}.parquet"
        print(f"Building strict labeled seed for {len(columns)} corrected variables", flush=True)
        run_checked([sys.executable, str(SCRIPTS / "08_build_custom_panel.py"), "--input", str(args.input),
                     "--output", str(seed), "--vars-file", str(variables_file), "--years", str(args.year),
                     "--dictionary", str(args.dictionary), "--codes", str(args.codes), "--column-lineage", str(args.lineage),
                     "--require-metadata", "--log-file", ""], proof / "parquet.log")
        seed_metadata = read_metadata(seed)
        expected = dataset.to_table(columns=[row["name"] for row in seed_metadata["variables"]], filter=ds.field("year") == args.year)
        receipt["formats"]["parquet"] = verify_representation(seed, expected)
        receipt["inputs"] = {"panel": {"path": str(args.input.resolve()), "sha256": seed_metadata["source_panel_sha256"]},
                             **{kind: {"path": seed_metadata["metadata_sources"][kind], "sha256": digest}
                                for kind, digest in seed_metadata["metadata_source_sha256"].items()},
                             "registry": {"path": str(args.registry.resolve()), "sha256": sha256_file(args.registry)},
                             "variable_map": {"path": str(args.variable_map.resolve()), "sha256": sha256_file(args.variable_map)}}
        for fmt in ("csv", "dta", "xlsx"):
            target = seed.with_suffix("." + fmt)
            print(f"Re-exporting and checking {fmt} from local embedded metadata", flush=True)
            run_checked([sys.executable, str(SCRIPTS / "08_build_custom_panel.py"), "--input", str(seed),
                         "--output", str(target), "--vars-file", str(variables_file), "--require-metadata", "--log-file", ""], proof / f"{fmt}.log")
            receipt["formats"][fmt] = verify_representation(target, expected)
            derived = read_metadata(target)
            if not any(parent.get("source_panel_sha256") == seed_metadata["source_panel_sha256"]
                       for parent in derived.get("upstream_export_provenance", [])):
                raise ValueError(f"{fmt} lost the original panel provenance")
        print("Checking the DTA in native Stata", flush=True)
        receipt["native_stata"] = verify_native_stata(seed.with_suffix(".dta"), expected, args.stata_binary, proof)
        lineage_records = [record for variable in seed_metadata["variables"] for record in variable.get("lineage_records", [])]
        compact_lineage = proof / "selected_lineage.parquet"
        pq.write_table(pa.Table.from_pylist(lineage_records), compact_lineage)
        stage09 = output / "panel_dictionary.xlsx"
        run_checked([sys.executable, str(SCRIPTS / "09_build_panel_dictionary.py"), "--input", str(seed),
                     "--dictionary", str(args.dictionary), "--codes", str(args.codes), "--column-lineage", str(compact_lineage),
                     "--output", str(stage09), "--require-ready"], proof / "stage09.log")
        stage09_metadata = read_metadata(stage09)
        if len(stage09_metadata["variables"]) != expected.num_columns:
            raise ValueError("Stage09 dictionary does not cover the exported schema")
        receipt["stage09"] = {"path": str(stage09), "variables": len(stage09_metadata["variables"]), "readiness_status": stage09_metadata["readiness_status"]}
        receipt["exporter_provenance"] = seed_metadata["exporter_provenance"]
        receipt["status"] = "pass"
    except Exception as exc:
        receipt["error"] = f"{type(exc).__name__}: {exc}"
        (proof / "failure.txt").write_text(traceback.format_exc())
        raise
    finally:
        receipt["artifacts"] = [{"path": str(path.resolve()), "sha256": sha256_file(path)}
                                for path in sorted(output.rglob("*")) if path.is_file() and path != output / "validation.json"]
        (output / "validation.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(f"PASS: {output / 'validation.json'}", flush=True)


if __name__ == "__main__":
    main()
