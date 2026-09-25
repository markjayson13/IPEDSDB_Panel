"""Round-trip checks for metadata-aware Stage 08 analyst exports.

These fixtures are intentionally independent of the external IPEDS data drive.
They check the exported artifact with its real reader, including the cases in
which a target format would otherwise silently change an observation.
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.csv as pcsv
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from helpers import load_script_module, run_script


@pytest.fixture
def export_fixture(tmp_path: Path) -> dict:
    root = tmp_path / "data_root"
    root.mkdir()
    panel = root / "panel.parquet"
    dictionary = root / "dictionary_lake.parquet"
    codes = root / "dictionary_codes.parquet"
    lineage = root / "column_lineage.parquet"
    table = pa.table({
        "year": pa.array([2023] * 6, type=pa.int32()),
        "UNITID": pa.array(range(100654, 100660), type=pa.int64()),
        "CONTROL": pa.array([1, 2, -1, None, 1, 2], type=pa.int64()),
        "INSTNM": pa.array(["Université 東京", "=SUM(1,2)", "#N/A", "00123", None, ""], type=pa.string()),
    })
    pq.write_table(table, panel)
    definitions = [
        {"year": 2023, "source_file": "HD", "varnumber": "00000001", "varname": "CONTROL",
         "varTitle": "Institution control", "longDescription": "Institution control classification.", "DataType": "N"},
        {"year": 2023, "source_file": "HD", "varnumber": "00000002", "varname": "INSTNM",
         "varTitle": "Institution name", "longDescription": "Name reported by the institution.", "DataType": "A"},
        {"year": 2022, "source_file": "HD", "varnumber": "00000001", "varname": "CONTROL",
         "varTitle": "Older control title", "longDescription": "Older definition.", "DataType": "N"},
        {"year": 2023, "source_file": "OTHER", "varnumber": "00000001", "varname": "CONTROL",
         "varTitle": "Unrelated control title", "longDescription": "Unrelated source definition.", "DataType": "N"},
    ]
    pq.write_table(pa.Table.from_pylist(definitions), dictionary)
    # Blank variable names deliberately exercise the source/year/number join.
    code_rows = [
        {"year": 2023, "source_file": "HD", "varnumber": "1", "varname": "", "codevalue": code,
         "valuelabel": label, "label_scope": "regular"}
        for code, label in [("1", "Public"), ("2", "Private"), ("-1", "Not reported")]
    ]
    code_rows.extend([
        {"year": 2022, "source_file": "HD", "varnumber": "1", "varname": "", "codevalue": "1",
         "valuelabel": "Older label", "label_scope": "regular"},
        {"year": 2023, "source_file": "OTHER", "varnumber": "1", "varname": "", "codevalue": "1",
         "valuelabel": "Wrong source label", "label_scope": "regular"},
    ])
    pq.write_table(pa.Table.from_pylist(code_rows), codes)
    pq.write_table(pa.Table.from_pylist([
        {"output_column": name, "source_varnames": name, "source_files": "HD", "lineage_kind": "direct_scalar"}
        for name in ("CONTROL", "INSTNM")
    ]), lineage)
    return {"root": root, "panel": panel, "dictionary": dictionary, "codes": codes,
            "lineage": lineage, "table": table}


def export_to(fixture: dict, output: Path, *extra: str):
    return run_script(
        "Scripts/08_build_custom_panel.py", "--input", fixture["panel"], "--output", output,
        "--vars", "CONTROL,INSTNM", "--dictionary", fixture["dictionary"], "--codes", fixture["codes"],
        "--column-lineage", fixture["lineage"], "--log-file", "", *extra,
        env={"IPEDSDB_ROOT": str(fixture["root"])},
    )


def read_metadata(output: Path) -> dict:
    metadata = json.loads(Path(str(output) + ".metadata.json").read_text(encoding="utf-8"))
    assert metadata["data_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert metadata["row_count"] >= 0
    for suffix in ("dictionary.csv", "value_labels.csv", "README.txt"):
        assert Path(str(output) + "." + suffix).is_file()
    return metadata


def test_stata_round_trip_embeds_labels_and_preserves_codes(export_fixture: dict) -> None:
    output = export_fixture["root"] / "extract.dta"
    result = export_to(export_fixture, output, "--require-metadata")
    assert result.returncode == 0, result.stdout

    with pd.read_stata(output, iterator=True, convert_categoricals=False) as reader:
        assert reader.variable_labels()["CONTROL"] == "Institution control"
        assert reader.variable_labels()["INSTNM"] == "Institution name"
        assert reader.value_labels()["CONTROL"] == {-1: "Not reported", 1: "Public", 2: "Private"}
        actual = reader.read()
    assert actual.columns.tolist() == ["year", "UNITID", "CONTROL", "INSTNM"]
    assert actual["CONTROL"].iloc[:3].tolist() == [1, 2, -1]
    assert pd.isna(actual.loc[3, "CONTROL"])
    assert actual["INSTNM"].tolist() == ["Université 東京", "=SUM(1,2)", "#N/A", "00123", "", ""]
    assert actual["UNITID"].tolist() == list(range(100654, 100660))
    metadata = read_metadata(output)
    assert metadata["years"] == [2023]
    assert metadata["row_count"] == 6
    assert any(issue["code"] == "stata_string_null" for issue in metadata["issues"])
    control = next(var for var in metadata["variables"] if var["name"] == "CONTROL")
    assert control["null_count"] == 1
    assert {record["source_file"] for record in control["source_metadata"]} == {"HD"}
    assert {record["year"] for record in control["value_label_records"]} == {2023}


def test_excel_round_trip_has_dictionary_and_literal_text(export_fixture: dict) -> None:
    output = export_fixture["root"] / "extract.xlsx"
    result = export_to(export_fixture, output, "--require-metadata")
    assert result.returncode == 0, result.stdout
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == ["data", "dictionary", "value_labels", "about", "issues"]
        data = workbook["data"]
        assert [cell.value for cell in data[1]] == ["year", "UNITID", "CONTROL", "INSTNM"]
        assert data["D2"].value == "Université 東京"
        for address, value in (("D3", "=SUM(1,2)"), ("D4", "#N/A"), ("D5", "00123")):
            assert data[address].value == value
            assert data[address].data_type == "s"
        assert data["C4"].value == -1
        assert data["C5"].value is None
        definitions = list(workbook["dictionary"].values)
        header = list(definitions[0])
        control = next(dict(zip(header, row)) for row in definitions[1:] if row[header.index("name")] == "CONTROL")
        assert control["label"] == "Institution control"
        assert control["source_files"] == "HD"
        labels = list(workbook["value_labels"].values)
        assert len(labels) == 4
    finally:
        workbook.close()
    read_metadata(output)


def test_csv_round_trip_preserves_null_empty_and_na_text(export_fixture: dict) -> None:
    fixture = export_fixture
    table = fixture["table"]
    names = pa.array([None, "", "NA", "00123", "=SUM(1,2)", "#N/A"], type=pa.string())
    table = table.set_column(table.schema.get_field_index("INSTNM"), "INSTNM", names)
    pq.write_table(table, fixture["panel"])
    output = fixture["root"] / "extract.csv"
    result = export_to(fixture, output, "--require-metadata")
    assert result.returncode == 0, result.stdout
    actual = pcsv.read_csv(output, convert_options=pcsv.ConvertOptions(
        column_types=table.schema, strings_can_be_null=True,
        quoted_strings_can_be_null=False, null_values=[""],
    ))
    assert actual.to_pydict() == table.to_pydict()
    metadata = read_metadata(output)
    assert next(var for var in metadata["variables"] if var["name"] == "INSTNM")["storage_type"] == "string"
    with Path(str(output) + ".value_labels.csv").open(newline="", encoding="utf-8") as handle:
        labels = list(csv.DictReader(handle))
    assert {(row["codevalue"], row["valuelabel"]) for row in labels} == {
        ("1", "Public"), ("2", "Private"), ("-1", "Not reported"),
    }


def test_parquet_round_trip_embeds_metadata_and_preserves_schema(export_fixture: dict) -> None:
    output = export_fixture["root"] / "extract.parquet"
    result = export_to(export_fixture, output, "--require-metadata")
    assert result.returncode == 0, result.stdout
    actual = pq.read_table(output)
    assert actual.to_pydict() == export_fixture["table"].to_pydict()
    assert actual.schema.equals(export_fixture["table"].schema, check_metadata=False)
    assert actual.schema.field("CONTROL").metadata[b"label"] == b"Institution control"
    embedded = json.loads(actual.schema.metadata[b"ipeds:export"])
    assert embedded["row_count"] == 6
    assert embedded["years"] == [2023]
    read_metadata(output)


@pytest.mark.parametrize("suffix,value", [("dta", 9007199254740993), ("xlsx", 1234567890123456)])
def test_numeric_precision_rejected_without_outputs(export_fixture: dict, suffix: str, value: int) -> None:
    table = export_fixture["table"]
    table = table.set_column(table.schema.get_field_index("CONTROL"), "CONTROL",
                             pa.array([value, None, 1, 2, -1, 1], type=pa.int64()))
    pq.write_table(table, export_fixture["panel"])
    output = export_fixture["root"] / f"unsafe.{suffix}"
    result = export_to(export_fixture, output)
    assert result.returncode != 0
    assert "precision" in result.stdout.lower() or "15 digits" in result.stdout.lower()
    assert not list(output.parent.glob(output.name + "*"))


def test_excel_oversized_text_rejected_without_outputs(export_fixture: dict) -> None:
    table = export_fixture["table"]
    table = table.set_column(table.schema.get_field_index("INSTNM"), "INSTNM",
                             pa.array(["x" * 32768, "", None, "A", "B", "C"], type=pa.string()))
    pq.write_table(table, export_fixture["panel"])
    output = export_fixture["root"] / "oversized.xlsx"
    result = export_to(export_fixture, output)
    assert result.returncode != 0
    assert "Excel cell text" in result.stdout
    assert not list(output.parent.glob(output.name + "*"))


@pytest.mark.parametrize("suffix", ["parquet", "csv", "dta", "xlsx"])
def test_empty_year_selection_still_writes_schema_and_sidecars(export_fixture: dict, suffix: str) -> None:
    output = export_fixture["root"] / f"empty.{suffix}"
    result = export_to(export_fixture, output, "--years", "1900")
    assert result.returncode == 0, result.stdout
    if suffix == "parquet":
        actual = pq.read_table(output)
        assert actual.num_rows == 0
        columns = actual.column_names
    elif suffix == "csv":
        with output.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
        assert len(rows) == 1
        columns = rows[0]
    elif suffix == "dta":
        actual = pd.read_stata(output, convert_categoricals=False)
        assert actual.empty
        columns = actual.columns.tolist()
    else:
        workbook = load_workbook(output)
        try:
            data_rows = list(workbook["data"].values)
            assert len(data_rows) == 1
            columns = list(data_rows[0])
        finally:
            workbook.close()
    assert columns == ["year", "UNITID", "CONTROL", "INSTNM"]
    assert read_metadata(output)["row_count"] == 0


@pytest.mark.parametrize("problem", ["missing", "conflicting"])
def test_required_metadata_failure_preserves_existing_export_package(export_fixture: dict, problem: str) -> None:
    output = export_fixture["root"] / "protected.csv"
    result = export_to(export_fixture, output, "--require-metadata")
    assert result.returncode == 0, result.stdout
    original = {path: path.read_bytes() for path in output.parent.glob(output.name + "*")}
    assert len(original) == 5
    definitions = pq.read_table(export_fixture["dictionary"]).to_pylist()
    if problem == "missing":
        definitions = [row for row in definitions if row["varname"] != "CONTROL"]
    else:
        conflicting = dict(definitions[0], varTitle="Conflicting current definition")
        definitions.append(conflicting)
    pq.write_table(pa.Table.from_pylist(definitions), export_fixture["dictionary"])
    result = export_to(export_fixture, output, "--require-metadata")
    assert result.returncode != 0
    assert "Metadata is incomplete" in result.stdout
    assert {path: path.read_bytes() for path in output.parent.glob(output.name + "*")} == original
    assert not list(output.parent.glob(".ipeds-export-*"))


def test_metadata_discovery_uses_input_root_and_infers_csv(export_fixture: dict, tmp_path: Path) -> None:
    root = export_fixture["root"]
    panels = root / "Panels"
    dictionary_dir = root / "Dictionary"
    lineage_dir = root / "Checks" / "wide_qc"
    for path in (panels, dictionary_dir, lineage_dir):
        path.mkdir(parents=True)
    panel = panels / "source.parquet"
    export_fixture["panel"].replace(panel)
    export_fixture["dictionary"].replace(dictionary_dir / "dictionary_lake.parquet")
    export_fixture["codes"].replace(dictionary_dir / "dictionary_codes.parquet")
    lineage_records = pq.read_table(export_fixture["lineage"]).to_pylist()
    with (lineage_dir / "qc_column_lineage.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(lineage_records[0]))
        writer.writeheader()
        writer.writerows(lineage_records)
    output = panels / "automatic.csv"
    result = run_script(
        "Scripts/08_build_custom_panel.py", "--input", panel, "--output", output,
        "--vars", "CONTROL,INSTNM", "--require-metadata", "--log-file", "",
        env={"IPEDSDB_ROOT": str(tmp_path / "unrelated_empty_root")},
    )
    assert result.returncode == 0, result.stdout
    metadata = read_metadata(output)
    assert metadata["format"] == "csv"
    assert metadata["metadata_status"] == "complete"
    assert Path(metadata["metadata_sources"]["dictionary"]) == dictionary_dir / "dictionary_lake.parquet"
    assert Path(metadata["metadata_sources"]["codes"]) == dictionary_dir / "dictionary_codes.parquet"
    with output.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    assert len(rows) == 7
    assert rows[0] == ["year", "UNITID", "CONTROL", "INSTNM"]


def test_portable_parquet_reexports_native_stata_labels_without_external_metadata(export_fixture: dict, tmp_path: Path) -> None:
    portable = export_fixture["root"] / "portable.parquet"
    result = export_to(export_fixture, portable, "--require-metadata")
    assert result.returncode == 0, result.stdout
    # Leave only the tagged Parquet: neither JSON/CSV sidecars nor source
    # dictionaries are available to the second invocation.
    for key in ("dictionary", "codes", "lineage"):
        export_fixture[key].unlink()
    for sidecar in portable.parent.glob(portable.name + ".*"):
        sidecar.unlink()
    output = export_fixture["root"] / "portable.dta"
    result = run_script(
        "Scripts/08_build_custom_panel.py", "--input", portable, "--output", output,
        "--vars", "CONTROL,INSTNM", "--require-metadata", "--log-file", "",
        env={"IPEDSDB_ROOT": str(tmp_path / "empty_root")},
    )
    assert result.returncode == 0, result.stdout
    with pd.read_stata(output, iterator=True, convert_categoricals=False) as reader:
        assert reader.variable_labels()["CONTROL"] == "Institution control"
        assert reader.value_labels()["CONTROL"] == {-1: "Not reported", 1: "Public", 2: "Private"}
        actual = reader.read()
    assert actual["UNITID"].tolist() == list(range(100654, 100660))
    assert actual["INSTNM"].iloc[0] == "Université 東京"
    metadata = read_metadata(output)
    assert metadata["metadata_status"] == "complete"
    assert metadata["metadata_source_modes"]["dictionary"] == "embedded_parquet"
    assert metadata["metadata_source_modes"]["codes"] == "embedded_parquet"
    assert metadata["metadata_sources"]["dictionary"] is None
    assert metadata["metadata_sources"]["codes"] is None


@pytest.mark.parametrize("problem", ["null_unitid", "invalid_unitid", "duplicate_key", "unknown_code"])
def test_readiness_checks_observations_without_recoding(export_fixture: dict, problem: str) -> None:
    table = export_fixture["table"]
    if problem == "null_unitid":
        values = pa.array([None, *range(100655, 100660)], type=pa.int64())
        column, expected = "UNITID", "invalid_panel_key"
    elif problem == "invalid_unitid":
        values = pa.array([100654.5, *range(100655, 100660)], type=pa.float64())
        column, expected = "UNITID", "invalid_panel_key"
    elif problem == "duplicate_key":
        values = pa.array([100654, 100655, 100654, *range(100657, 100660)], type=pa.int64())
        column, expected = "UNITID", "duplicate_panel_key"
    else:
        values = pa.array([99, 2, -1, None, 1, 2], type=pa.int64())
        column, expected = "CONTROL", "unknown_observed_code"
    table = table.set_column(table.schema.get_field_index(column), column, values)
    pq.write_table(table, export_fixture["panel"])
    output = export_fixture["root"] / "not_ready.parquet"
    strict = export_to(export_fixture, output, "--require-ready", "--batch-rows", "2")
    assert strict.returncode != 0
    assert "observations are not ready" in strict.stdout
    assert not list(output.parent.glob(output.name + "*"))
    diagnostic = export_to(export_fixture, output, "--batch-rows", "2")
    assert diagnostic.returncode == 0, diagnostic.stdout
    metadata = read_metadata(output)
    assert metadata["metadata_status"] == "complete"
    assert metadata["readiness_status"] == "incomplete"
    assert metadata["observation_validation"]["status"] == "incomplete"
    assert any(issue["code"] == expected for issue in metadata["issues"])
    assert pq.read_table(output).to_pydict() == table.to_pydict()


def test_observed_codes_are_checked_in_their_reporting_year(export_fixture: dict) -> None:
    table = export_fixture["table"]
    table = table.set_column(0, "year", pa.array([2022, 2023, 2023, 2023, 2023, 2023], type=pa.int32()))
    # Code 2 exists in the 2023 codebook but not the 2022 codebook.
    table = table.set_column(2, "CONTROL", pa.array([2, 2, -1, None, 1, 2], type=pa.int64()))
    pq.write_table(table, export_fixture["panel"])
    output = export_fixture["root"] / "year_scoped.csv"
    result = export_to(export_fixture, output)
    assert result.returncode == 0, result.stdout
    metadata = read_metadata(output)
    issue = next(issue for issue in metadata["issues"] if issue["code"] == "unknown_observed_code")
    assert issue["count"] == 1
    assert issue["examples"] == [{"year": 2022, "value": "2"}]


def test_export_records_source_and_exporter_fingerprints(export_fixture: dict) -> None:
    output = export_fixture["root"] / "fingerprinted.csv"
    result = export_to(export_fixture, output, "--require-ready")
    assert result.returncode == 0, result.stdout
    metadata = read_metadata(output)
    assert metadata["source_panel_sha256"] == hashlib.sha256(export_fixture["panel"].read_bytes()).hexdigest()
    assert metadata["source_fingerprint"]["kind"] == "file"
    assert metadata["readiness_status"] == "complete"
    assert "Scripts/export_integrity.py" in metadata["exporter_provenance"]["files"]


def test_source_mutation_with_same_row_count_preserves_prior_package(export_fixture: dict, monkeypatch) -> None:
    output = export_fixture["root"] / "protected.parquet"
    result = export_to(export_fixture, output, "--require-ready")
    assert result.returncode == 0, result.stdout
    before = {path: path.read_bytes() for path in output.parent.glob(output.name + "*")}
    module = load_script_module("stage08_source_mutation", "Scripts/08_build_custom_panel.py")
    original_write = module.write_stream
    def mutate_source_after_write(*args, **kwargs):
        rows = original_write(*args, **kwargs)
        source = export_fixture["table"].set_column(2, "CONTROL", pa.array([2, 2, -1, None, 1, 2], type=pa.int64()))
        pq.write_table(source, export_fixture["panel"])
        return rows
    monkeypatch.setattr(module, "write_stream", mutate_source_after_write)
    monkeypatch.setattr("sys.argv", [
        "08_build_custom_panel.py", "--input", str(export_fixture["panel"]), "--output", str(output),
        "--vars", "CONTROL,INSTNM", "--dictionary", str(export_fixture["dictionary"]),
        "--codes", str(export_fixture["codes"]), "--column-lineage", str(export_fixture["lineage"]),
        "--log-file", "", "--require-ready",
    ])
    with pytest.raises(ValueError, match="Source changed during export"):
        module.main()
    assert {path: path.read_bytes() for path in output.parent.glob(output.name + "*")} == before
